"""Small, deterministic document fixtures for import parser tests."""

from pathlib import Path


DOCX_PARAGRAPH_BEFORE = "表格之前 α"
DOCX_TABLE_ROWS = (("项目", "青禾-47"), ("负责人", "林岚"))
DOCX_PARAGRAPH_AFTER = "表格之后 β"


def _pdf_string(value):
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _write_pdf(path, objects):
    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    Path(path).write_bytes(payload)


def write_text_pdf(path, pages):
    """Write a real-offset PDF with one Helvetica text stream per page."""
    pages = tuple(pages)
    first_page_id = 4
    page_ids = [first_page_id + index * 2 for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for index, value in enumerate(pages):
        page_id = page_ids[index]
        content_id = page_id + 1
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
            ).encode("ascii")
        )
        if value is None:
            content = b"q Q"
        else:
            rendered = _pdf_string(str(value))
            content = f"BT /F1 12 Tf 72 720 Td ({rendered}) Tj ET".encode("latin-1")
        objects.append(
            f"<< /Length {len(content)} >>\nstream\n".encode("ascii")
            + content
            + b"\nendstream"
        )
    _write_pdf(path, objects)


def write_image_pdf(path):
    """Write a one-page PDF containing an image XObject and no text."""
    image = b"\x00"
    content = b"q 1 0 0 1 0 0 cm /Im0 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R] /Count 1 >>",
        (
            b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 "
            b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\nstream\n"
            + image
            + b"\nendstream"
        ),
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 1 1] "
            b"/Resources << /XObject << /Im0 3 0 R >> >> /Contents 5 0 R >>"
        ),
        f"<< /Length {len(content)} >>\nstream\n".encode("ascii")
        + content
        + b"\nendstream",
    ]
    _write_pdf(path, objects)


def write_nested_form_pdf(path):
    """Write two pages sharing nested Forms, a resource cycle, and an Image."""
    image = b"\x00" * 128
    inner_content = b"BT /F1 12 Tf 10 20 Td (Nested form text) Tj ET"
    outer_content = b"q /Inner Do /Im0 Do Q"
    page_content = b"q /Outer Do /Im0 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [7 0 R 9 0 R] /Count 2 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            b"<< /Type /XObject /Subtype /Image /Width 128 /Height 1 "
            b"/ColorSpace /DeviceGray /BitsPerComponent 8 /Length 128 >>\nstream\n"
            + image
            + b"\nendstream"
        ),
        (
            b"<< /Type /XObject /Subtype /Form /FormType 1 /BBox [0 0 200 50] "
            b"/Resources << /Font << /F1 3 0 R >> /XObject << /Outer 6 0 R >> >> "
            + f"/Length {len(inner_content)} >>\nstream\n".encode("ascii")
            + inner_content
            + b"\nendstream"
        ),
        (
            b"<< /Type /XObject /Subtype /Form /FormType 1 /BBox [0 0 200 50] "
            b"/Resources << /XObject << /Inner 5 0 R /Im0 4 0 R >> >> "
            + f"/Length {len(outer_content)} >>\nstream\n".encode("ascii")
            + outer_content
            + b"\nendstream"
        ),
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /XObject << /Outer 6 0 R /Im0 4 0 R >> >> /Contents 8 0 R >>",
        f"<< /Length {len(page_content)} >>\nstream\n".encode("ascii")
        + page_content
        + b"\nendstream",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /XObject << /Outer 6 0 R /Im0 4 0 R >> >> /Contents 10 0 R >>",
        f"<< /Length {len(page_content)} >>\nstream\n".encode("ascii")
        + page_content
        + b"\nendstream",
    ]
    _write_pdf(path, objects)
    return {
        "page_content_bytes": len(page_content) * 2,
        "form_content_bytes": len(inner_content) + len(outer_content),
        "image_content_bytes": len(image),
    }


def write_docx(path):
    """Write paragraph -> table -> paragraph using python-docx."""
    from docx import Document

    document = Document()
    document.add_paragraph(DOCX_PARAGRAPH_BEFORE)
    table = document.add_table(rows=len(DOCX_TABLE_ROWS), cols=2)
    for row, values in zip(table.rows, DOCX_TABLE_ROWS):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph(DOCX_PARAGRAPH_AFTER)
    document.save(path)
