"""Deterministic, bounded parsers for locally imported documents."""

from dataclasses import dataclass
import io
import math
from pathlib import Path
import unicodedata
import zipfile


SUPPORTED_EXTENSIONS = frozenset({".md", ".txt", ".pdf", ".docx"})


@dataclass(frozen=True)
class ImportLimits:
    max_file_bytes: int = 20 * 1024 * 1024
    max_extracted_bytes: int = 512 * 1024
    max_pdf_pages: int = 500
    max_pdf_content_bytes: int = 64 * 1024 * 1024
    max_docx_entries: int = 2000
    max_docx_uncompressed_bytes: int = 100 * 1024 * 1024
    max_docx_ratio: int = 100

    def __post_init__(self):
        integer_fields = (
            "max_file_bytes",
            "max_extracted_bytes",
            "max_pdf_pages",
            "max_pdf_content_bytes",
            "max_docx_entries",
            "max_docx_uncompressed_bytes",
        )
        for name in integer_fields:
            value = getattr(self, name)
            if type(value) is not int or value < 0:
                raise ValueError("invalid import limit")
        ratio = self.max_docx_ratio
        if (isinstance(ratio, bool) or not isinstance(ratio, (int, float))
                or not math.isfinite(ratio) or ratio < 0):
            raise ValueError("invalid import limit")


@dataclass(frozen=True)
class ParsedUnit:
    text: str
    location: dict


@dataclass(frozen=True)
class ParsedDocument:
    logical_path: str
    units: tuple[ParsedUnit, ...]
    warnings: tuple[dict, ...]
    stats: dict


class DocumentParseError(ValueError):
    def __init__(self, code):
        self.code = code
        super().__init__(code)


def _fail(code):
    raise DocumentParseError(code)


def _read_source(source, limit):
    try:
        with source.open("rb") as stream:
            raw = stream.read(limit + 1)
    except (OSError, ValueError, TypeError):
        _fail("DOCUMENT_CORRUPT")
    if len(raw) > limit:
        _fail("DOCUMENT_LIMIT_EXCEEDED")
    return raw


def _check_extracted_size(size, limits):
    if size > limits.max_extracted_bytes:
        _fail("DOCUMENT_LIMIT_EXCEEDED")


def _common_stats(extension, source_bytes, extracted_bytes, units):
    return {
        "format": extension[1:],
        "source_bytes": source_bytes,
        "extracted_bytes": extracted_bytes,
        "unit_count": len(units),
    }


def _parse_text(raw, logical_path, extension, limits):
    try:
        text = raw.decode("utf-8", errors="strict")
    except UnicodeDecodeError:
        _fail("TEXT_INVALID_UTF8")
    if any(unicodedata.category(character) == "Cc" and character not in "\t\n\r"
           for character in text):
        _fail("TEXT_INVALID_CHARACTER")
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    extracted_bytes = len(normalized.encode("utf-8"))
    _check_extracted_size(extracted_bytes, limits)
    if not normalized:
        lines = []
    else:
        parts = normalized.split("\n")
        lines = [part + "\n" for part in parts[:-1]]
        if parts[-1]:
            lines.append(parts[-1])
    units = tuple(
        ParsedUnit(line, {"kind": "text_lines", "start": number, "end": number})
        for number, line in enumerate(lines, 1)
    )
    return ParsedDocument(
        logical_path,
        units,
        (),
        _common_stats(extension, len(raw), extracted_bytes, units),
    )


def _pdf_library():
    try:
        from pypdf import PdfReader
    except Exception:
        _fail("DOCUMENT_PARSER_UNAVAILABLE")
    return PdfReader


def _decoded_pdf_content(page):
    contents = page.get_contents()
    if contents is None:
        return b""
    return contents.get_data()


def _pdf_object(value):
    get_object = getattr(value, "get_object", None)
    return get_object() if get_object is not None else value


def _pdf_object_identity(reference, value):
    indirect = reference if hasattr(reference, "idnum") else getattr(
        value, "indirect_reference", None
    )
    if indirect is not None and hasattr(indirect, "idnum"):
        return (indirect.idnum, indirect.generation)
    return ("direct", id(value))


def _pdf_xobjects(owner):
    resources = _pdf_object(owner.get("/Resources"))
    if resources is None:
        return ()
    xobjects = _pdf_object(resources.get("/XObject"))
    return tuple(xobjects.values()) if xobjects is not None else ()


def _add_reachable_form_content(page, seen, content_bytes, limits):
    pending = list(_pdf_xobjects(page))
    while pending:
        reference = pending.pop()
        value = _pdf_object(reference)
        if value.get("/Subtype") != "/Form":
            continue
        identity = _pdf_object_identity(reference, value)
        if identity in seen:
            continue
        seen.add(identity)
        content_bytes += len(value.get_data())
        if content_bytes > limits.max_pdf_content_bytes:
            _fail("DOCUMENT_LIMIT_EXCEEDED")
        pending.extend(_pdf_xobjects(value))
    return content_bytes


def _parse_pdf(raw, logical_path, extension, limits):
    PdfReader = _pdf_library()
    try:
        reader = PdfReader(io.BytesIO(raw), strict=True)
        if reader.is_encrypted:
            _fail("PDF_ENCRYPTED")
        page_count = len(reader.pages)
        if page_count > limits.max_pdf_pages:
            _fail("DOCUMENT_LIMIT_EXCEEDED")
        units = []
        warnings = []
        content_bytes = 0
        extracted_bytes = 0
        seen_forms = set()
        for page_number, page in enumerate(reader.pages, 1):
            content_bytes += len(_decoded_pdf_content(page))
            if content_bytes > limits.max_pdf_content_bytes:
                _fail("DOCUMENT_LIMIT_EXCEEDED")
            content_bytes = _add_reachable_form_content(
                page, seen_forms, content_bytes, limits
            )
            text = page.extract_text()
            if not text or not text.strip():
                warnings.append({"code": "PDF_PAGE_TEXT_NOT_FOUND", "page": page_number})
                continue
            encoded_size = len(text.encode("utf-8"))
            extracted_bytes += encoded_size
            _check_extracted_size(extracted_bytes, limits)
            units.append(ParsedUnit(text, {"kind": "pdf_page", "page": page_number}))
        if not units:
            _fail("PDF_TEXT_NOT_FOUND")
    except DocumentParseError:
        raise
    except MemoryError:
        _fail("DOCUMENT_LIMIT_EXCEEDED")
    except Exception:
        _fail("DOCUMENT_CORRUPT")
    stats = _common_stats(extension, len(raw), extracted_bytes, units)
    stats.update({"page_count": page_count, "pdf_content_bytes": content_bytes})
    return ParsedDocument(logical_path, tuple(units), tuple(warnings), stats)


def _inspect_docx_archive(raw, limits):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_docx_entries:
                _fail("DOCUMENT_LIMIT_EXCEEDED")
            uncompressed_bytes = 0
            for entry in entries:
                uncompressed_bytes += entry.file_size
                if uncompressed_bytes > limits.max_docx_uncompressed_bytes:
                    _fail("DOCUMENT_LIMIT_EXCEEDED")
                if entry.compress_size == 0:
                    if entry.file_size:
                        _fail("DOCUMENT_LIMIT_EXCEEDED")
                elif entry.file_size / entry.compress_size > limits.max_docx_ratio:
                    _fail("DOCUMENT_LIMIT_EXCEEDED")
            if archive.testzip() is not None:
                _fail("DOCUMENT_CORRUPT")
    except DocumentParseError:
        raise
    except MemoryError:
        _fail("DOCUMENT_LIMIT_EXCEEDED")
    except (OSError, ValueError, RuntimeError, zipfile.BadZipFile, zipfile.LargeZipFile):
        _fail("DOCUMENT_CORRUPT")
    return len(entries), uncompressed_bytes


def _docx_library():
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except Exception:
        _fail("DOCUMENT_PARSER_UNAVAILABLE")
    return Document, Paragraph, Table


def _parse_docx(raw, logical_path, extension, limits):
    entry_count, uncompressed_bytes = _inspect_docx_archive(raw, limits)
    Document, Paragraph, Table = _docx_library()
    try:
        document = Document(io.BytesIO(raw))
        units = []
        paragraph_index = 0
        table_index = 0
        extracted_bytes = 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph_index += 1
                text = block.text
                location = {"kind": "docx_paragraph", "paragraph": paragraph_index}
                extracted_bytes += len(text.encode("utf-8"))
                _check_extracted_size(extracted_bytes, limits)
                units.append(ParsedUnit(text, location))
            elif isinstance(block, Table):
                table_index += 1
                for row_index, row in enumerate(block.rows, 1):
                    text = "\t".join(cell.text for cell in row.cells)
                    location = {
                        "kind": "docx_table_row",
                        "table": table_index,
                        "row": row_index,
                    }
                    extracted_bytes += len(text.encode("utf-8"))
                    _check_extracted_size(extracted_bytes, limits)
                    units.append(ParsedUnit(text, location))
    except DocumentParseError:
        raise
    except MemoryError:
        _fail("DOCUMENT_LIMIT_EXCEEDED")
    except Exception:
        _fail("DOCUMENT_CORRUPT")
    stats = _common_stats(extension, len(raw), extracted_bytes, units)
    stats.update({
        "docx_entries": entry_count,
        "docx_uncompressed_bytes": uncompressed_bytes,
    })
    return ParsedDocument(logical_path, tuple(units), (), stats)


def parse_document(source: Path, logical_path: str, limits: ImportLimits = ImportLimits()):
    """Parse a supported local document without returning unstable library errors."""
    if not isinstance(source, Path) or not isinstance(logical_path, str):
        _fail("DOCUMENT_CORRUPT")
    extension = Path(logical_path).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        _fail("IMPORT_FORMAT_UNSUPPORTED")
    if not isinstance(limits, ImportLimits):
        _fail("DOCUMENT_CORRUPT")
    raw = _read_source(source, limits.max_file_bytes)
    if extension in {".md", ".txt"}:
        return _parse_text(raw, logical_path, extension, limits)
    if extension == ".pdf":
        return _parse_pdf(raw, logical_path, extension, limits)
    return _parse_docx(raw, logical_path, extension, limits)
