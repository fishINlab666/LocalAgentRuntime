"""Fixed two-run acceptance for managed local document imports."""

import copy
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import tempfile
import threading
import uuid

from .agents import build_file_engine, builtin_agent
from .approvals import RunControl
from .directory_evaluation import context_checks
from .imports import ImportStore
from .provider import ProviderError
from .session_store import SessionStore
from .sessions import RunSubmission, SessionService
from .state_backup import StateBackup
from .trace import Trace


MAX_MODEL_CALLS = 12
EXPECTED_RUNS = 2
_FIRST_FACTS = ("Budget 42", "Owner Mei")
_HIDDEN_FACTS = (*_FIRST_FACTS, "TXT-READY", "MD-ALPHA")
_MANDATORY_CHECKS = frozenset({
    "mixed_formats_and_ignored_file",
    "initial_request_has_no_document_body",
    "first_run_completed",
    "first_run_search_then_reads",
    "first_run_tool_results_round_trip",
    "pdf_page_location",
    "docx_table_location",
    "first_run_facts",
    "other_import_isolated",
    "external_originals_removed",
    "restart_keeps_managed_session",
    "restart_can_read_managed_copy",
    "restore_preserves_logical_identity",
    "restore_rebinds_file_identity",
    "restored_follow_up_completed",
    "history_projection_present",
    "follow_up_search_then_read",
    "follow_up_tool_results_round_trip",
    "text_source_location",
    "follow_up_facts",
    "other_import_secret_never_reached_provider",
    "fixed_run_and_call_budget",
})


def _pdf_string(value):
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _write_text_pdf(path, pages):
    pages = tuple(pages)
    page_ids = [4 + index * 2 for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, value in zip(page_ids, pages):
        content_id = page_id + 1
        objects.append((
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode("ascii"))
        content = (
            f"BT /F1 12 Tf 72 720 Td ({_pdf_string(value)}) Tj ET"
        ).encode("latin-1")
        objects.append(
            f"<< /Length {len(content)} >>\nstream\n".encode("ascii")
            + content + b"\nendstream"
        )
    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    Path(path).write_bytes(payload)


def _write_docx(path):
    from docx import Document

    document = Document()
    document.add_paragraph("Acceptance owner record")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (
        ("Acceptance fact", "Owner Mei"),
        ("State", "Ready"),
    )):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("End of owner record")
    document.save(path)


def _fixture(source):
    source.mkdir(parents=True, mode=0o700)
    paths = {
        "资料/overview.md": source / "overview.md",
        "资料/checkpoint.txt": source / "checkpoint.txt",
        "资料/budget.pdf": source / "budget.pdf",
        "资料/owner.docx": source / "owner.docx",
        "资料/preview.png": source / "preview.png",
    }
    paths["资料/overview.md"].write_text(
        "# Acceptance packet\nMD marker: MD-ALPHA\n", encoding="utf-8"
    )
    paths["资料/checkpoint.txt"].write_text(
        "Release checkpoint: TXT-READY\n", encoding="utf-8"
    )
    _write_text_pdf(paths["资料/budget.pdf"], (
        "Acceptance cover", "Acceptance fact: Budget 42",
    ))
    _write_docx(paths["资料/owner.docx"])
    paths["资料/preview.png"].write_bytes(b"\x89PNG\r\n\x1a\nignored")
    return {name: path.read_bytes() for name, path in paths.items()}


class _RecordingProvider:
    def __init__(self, provider):
        self.provider = provider
        self.metadata = copy.deepcopy(provider.metadata)
        self.requests = []

    def complete(self, messages, tools, timeout):
        if len(self.requests) >= MAX_MODEL_CALLS:
            raise ProviderError("MODEL_CALL_LIMIT")
        self.requests.append({
            "messages": copy.deepcopy(messages),
            "tools": copy.deepcopy(tools),
        })
        return self.provider.complete(messages, tools, timeout)


def _open_service(path):
    store = SessionStore.open(path)
    try:
        store.recover_interrupted(uuid.uuid4().hex)
        service = SessionService(store)
        service.recover_imports()
        return store, service
    except BaseException:
        store.close()
        raise


def _publish(service, entries, *, name, ignored=()):
    supported = [(logical_path, raw) for logical_path, raw in entries
                 if Path(logical_path).suffix.casefold() in {".md", ".txt", ".pdf", ".docx"}]
    request = {
        "kind": "folder" if len(supported) + len(ignored) > 1 else "file",
        "name": name,
        "agent_id": "directory-qa",
        "files": [
            {"logical_path": logical_path, "bytes": len(raw)}
            for logical_path, raw in supported
        ],
        "ignored": [
            {"logical_path": logical_path, "bytes": len(raw)}
            for logical_path, raw in ignored
        ],
    }
    imports = service.imports
    batch = imports.begin(request)
    for slot, (_logical_path, raw) in zip(batch.files, supported):
        imports.add_file(batch.id, slot.slot_id, BytesIO(raw), len(raw))
    published = imports.start_finalize(batch.id).run()
    return published, service.attach_import(published, request), request


def _chunk_for_extension(published, extension):
    record = next(item for item in published.file_records
                  if item["extension"] == extension)
    prefix = f"documents/{record['source_id']}/"
    return next(path.relative_to(published.workspace).as_posix()
                for path in published.chunk_paths
                if path.relative_to(published.workspace).as_posix().startswith(prefix))


def _events(path):
    return [json.loads(line) for line in Path(path).read_text(encoding="utf-8").splitlines()]


def _tool_requests(events):
    result = []
    for event in events:
        if event.get("event") != "tool.requested":
            continue
        data = event["data"]
        try:
            arguments = json.loads(data["arguments"])
        except (KeyError, TypeError, ValueError, RecursionError):
            arguments = {}
        result.append({"id": data.get("id"), "name": data.get("name"),
                       "arguments": arguments})
    return result


def _ordered_reads(calls, paths):
    search_index = next((index for index, item in enumerate(calls)
                         if item["name"] == "search_documents"), None)
    if search_index is None:
        return False
    found = {}
    for index, item in enumerate(calls):
        path = item["arguments"].get("path")
        if item["name"] == "read_file" and path in paths and path not in found:
            found[path] = index
    return set(found) == set(paths) and all(index > search_index for index in found.values())


def _source_locations(answer, logical_path):
    for citation in (answer or {}).get("citations", []):
        source = citation.get("source", {})
        if source.get("logical_path") == logical_path:
            return source.get("locations")
    return None


def _run(service, session, provider, output, question, request_id, cancel):
    prepared = service.submit(session.id, RunSubmission(
        request_id, question, "files", session.scope, None, None, {}
    ))
    trace = Trace(output / "traces", service.trace_root(session),
                  debug_content=True, run_id=prepared.run_id)
    result = service.execute(
        prepared, provider, trace, control=RunControl(cancel, run_timeout=120)
    )
    return result, trace


def evaluate_imports(provider_factory, directory: Path,
                     cancel: threading.Event | None = None) -> dict:
    """Run one bounded synthetic import session and preserve its raw evidence."""
    cancel = cancel or threading.Event()
    output = Path(directory).resolve() / ("import-evaluation-" + uuid.uuid4().hex)
    output.mkdir(parents=True, mode=0o700)
    report = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "gate": "NOT_RUN",
        "automatic_checks_passed": False,
        "semantic_review": "pending",
        "expected_runs": EXPECTED_RUNS,
        "completed_runs": 0,
        "model_calls": 0,
        "max_model_calls": MAX_MODEL_CALLS,
        "runs": [],
        "checks": {},
        "report_path": str(output / "report.json"),
    }

    def save():
        with Path(report["report_path"]).open("x", encoding="utf-8") as file:
            json.dump(report, file, ensure_ascii=False, indent=2)
        return report

    try:
        provider = _RecordingProvider(provider_factory())
    except ProviderError as error:
        report["error"] = error.code
        return save()
    report["provider"] = provider.metadata

    store = None
    try:
        with tempfile.TemporaryDirectory(prefix="local-agent-import-eval-") as temporary:
            root = Path(temporary)
            state = root / "state"
            source = root / "external-source"
            fixture = _fixture(source)
            entries = [(name, raw) for name, raw in fixture.items()
                       if not name.endswith(".png")]
            ignored = [(name, raw) for name, raw in fixture.items()
                       if name.endswith(".png")]

            store, service = _open_service(state)
            published, session, request = _publish(
                service, entries, name="混合资料验收", ignored=ignored
            )
            pdf_path = _chunk_for_extension(published, ".pdf")
            docx_path = _chunk_for_extension(published, ".docx")
            text_path = _chunk_for_extension(published, ".txt")

            first_start = len(provider.requests)
            first, first_trace = _run(
                service,
                session,
                provider,
                output,
                "先用 search_documents 定位资料，再逐字核对预算和负责人并引用原文位置。",
                "import-acceptance-first",
                cancel,
            )
            first_requests = provider.requests[first_start:]
            first_events = _events(first_trace.path)
            first_calls = _tool_requests(first_events)
            report["runs"].append(first)
            initial = json.dumps(first_requests[0], ensure_ascii=False) if first_requests else ""
            report["checks"].update({
                "mixed_formats_and_ignored_file": (
                    [Path(item["logical_path"]).suffix for item in request["files"]]
                    == [".md", ".txt", ".pdf", ".docx"]
                    and [item["logical_path"] for item in request["ignored"]]
                    == ["资料/preview.png"]
                ),
                "initial_request_has_no_document_body": (
                    bool(first_requests)
                    and all(fact not in initial for fact in _HIDDEN_FACTS)
                    and str(source) not in initial
                ),
                "first_run_completed": first.get("state") == "completed",
                "first_run_search_then_reads": _ordered_reads(
                    first_calls, (pdf_path, docx_path)
                ),
                "first_run_tool_results_round_trip": context_checks(
                    first_events
                )["tool_results_in_next_context"],
                "pdf_page_location": _source_locations(
                    first.get("answer"), "资料/budget.pdf"
                ) == [{"kind": "pdf_page", "page": 2}],
                "docx_table_location": _source_locations(
                    first.get("answer"), "资料/owner.docx"
                ) == [{"kind": "docx_table_row", "table": 1, "row": 1}],
                "first_run_facts": all(
                    fact in (first.get("answer") or {}).get("answer", "")
                    for fact in _FIRST_FACTS
                ),
            })
            if first.get("state") != "completed" or cancel.is_set():
                raise RuntimeError("FIRST_RUN_FAILED")

            secret = b"OTHER-IMPORT-SECRET\n"
            other, _other_session, _ = _publish(
                service,
                [("private/other.txt", secret)],
                name="另一批资料",
            )
            mapper = service.resolver.resolve(session).source_mapper
            other_path = _chunk_for_extension(other, ".txt")
            engine = build_file_engine(
                builtin_agent("directory"), mapper.workspace, None,
                source_mapper=mapper,
            )
            isolated_search = engine.registry.get("search_documents").execute({
                "query": "OTHER-IMPORT-SECRET",
            })
            isolated_read = engine.registry.get("read_file").execute({
                "path": other_path,
            })
            report["checks"]["other_import_isolated"] = (
                isolated_search == {
                    "ok": True,
                    "query": "OTHER-IMPORT-SECRET",
                    "evidence_role": "catalog",
                    "matches": [],
                }
                and isolated_read.get("error", {}).get("code") == "PATH_DENIED"
            )

            original_session = service.load(session.id)
            original_resolved = service.resolver.resolve(original_session)
            original_read_identity = original_resolved.read_identity
            original_write_identity = original_resolved.write_identity
            shutil.rmtree(source)
            report["checks"]["external_originals_removed"] = not source.exists()

            store.close()
            store = None
            store, service = _open_service(state)
            reopened = service.load(session.id)
            reopened_resolved = service.resolver.resolve(reopened)
            report["checks"]["restart_keeps_managed_session"] = (
                reopened.id == original_session.id
                and reopened.import_id == original_session.import_id
                and reopened.agent_snapshot == original_session.agent_snapshot
                and reopened_resolved.source_mapper is not None
            )
            reopened_engine = build_file_engine(
                builtin_agent("directory"), reopened_resolved.read_root, None,
                source_mapper=reopened_resolved.source_mapper,
            )
            probe_search = reopened_engine.registry.get("search_documents").execute({
                "query": "Acceptance fact",
            })
            probe_read = reopened_engine.registry.get("read_file").execute({
                "path": pdf_path,
            })
            reopened_engine.policy.accept(
                "read_file", {"path": pdf_path}, probe_read
            )
            probe_answer = reopened_engine.policy.validate(json.dumps({
                "status": "answered",
                "answer": "重启后仍可读取 Budget 42。",
                "citations": [{"path": pdf_path, "start_line": 2, "end_line": 2}],
            }, ensure_ascii=False))
            report["checks"]["restart_can_read_managed_copy"] = (
                probe_search.get("ok") is True
                and any(item.get("path") == pdf_path
                        for item in probe_search.get("matches", []))
                and probe_read.get("ok") is True
                and "Budget 42" in probe_read.get("content", "")
                and _source_locations(
                    probe_answer, "资料/budget.pdf"
                ) == [{"kind": "pdf_page", "page": 2}]
            )

            backup = service.backup(output / "backup")
            store.close()
            store = None
            restored_state = root / "restored-state"
            StateBackup.restore_bundle(
                Path(backup["database_path"]),
                Path(backup["sidecar_path"]),
                restored_state,
            )
            store, service = _open_service(restored_state)
            restored = service.load(session.id)
            restored_resolved = service.resolver.resolve(restored)
            read_details = os.stat(restored_resolved.read_root)
            write_details = os.stat(restored_resolved.write_root)
            report["checks"]["restore_preserves_logical_identity"] = (
                restored.id == original_session.id
                and restored.import_id == original_session.import_id
                and restored.agent_id == original_session.agent_id
                and restored.agent_revision == original_session.agent_revision
                and restored.agent_snapshot == original_session.agent_snapshot
            )
            report["checks"]["restore_rebinds_file_identity"] = (
                restored_resolved.read_identity == (read_details.st_dev, read_details.st_ino)
                and restored_resolved.write_identity == (write_details.st_dev, write_details.st_ino)
                and restored_resolved.read_identity != original_read_identity
                and restored_resolved.write_identity != original_write_identity
                and restored_resolved.read_root.is_relative_to(restored_state.resolve())
                and restored_resolved.write_root.is_relative_to(restored_state.resolve())
            )

            second_start = len(provider.requests)
            second, second_trace = _run(
                service,
                restored,
                provider,
                output,
                ("沿用上一轮结论，再用 search_documents 核对 Release checkpoint；"
                 "保留标识原文并引用依据。"),
                "import-acceptance-follow-up",
                cancel,
            )
            second_requests = provider.requests[second_start:]
            second_events = _events(second_trace.path)
            second_calls = _tool_requests(second_events)
            report["runs"].append(second)
            second_initial = second_requests[0]["messages"] if second_requests else []
            history = []
            for message in second_initial:
                if message.get("role") != "user":
                    continue
                try:
                    payload = json.loads(message.get("content", ""))
                except (TypeError, ValueError, RecursionError):
                    continue
                if payload.get("source_kind") == "session_history_projection":
                    history.append(payload)
            history_json = json.dumps(history, ensure_ascii=False)
            report["checks"].update({
                "restored_follow_up_completed": second.get("state") == "completed",
                "history_projection_present": (
                    len(history) == 1
                    and all(fact in history_json for fact in _FIRST_FACTS)
                    and not any(message.get("role") == "tool" for message in second_initial)
                ),
                "follow_up_search_then_read": _ordered_reads(second_calls, (text_path,)),
                "follow_up_tool_results_round_trip": context_checks(
                    second_events
                )["tool_results_in_next_context"],
                "text_source_location": _source_locations(
                    second.get("answer"), "资料/checkpoint.txt"
                ) == [{"kind": "text_lines", "start": 1, "end": 1}],
                "follow_up_facts": all(
                    fact in (second.get("answer") or {}).get("answer", "")
                    for fact in (*_FIRST_FACTS, "TXT-READY")
                ),
                "other_import_secret_never_reached_provider": all(
                    "OTHER-IMPORT-SECRET" not in json.dumps(item, ensure_ascii=False)
                    for item in provider.requests
                ),
            })
    except Exception as error:
        report["error"] = getattr(error, "code", type(error).__name__)
    finally:
        if store is not None:
            store.close()

    report["model_calls"] = len(provider.requests)
    report["completed_runs"] = sum(
        result.get("state") == "completed" for result in report["runs"]
    )
    report["checks"]["fixed_run_and_call_budget"] = (
        len(report["runs"]) == EXPECTED_RUNS
        and report["completed_runs"] == EXPECTED_RUNS
        and report["model_calls"] <= MAX_MODEL_CALLS
    )
    passed = (
        "error" not in report
        and set(report["checks"]) == _MANDATORY_CHECKS
        and all(report["checks"].values())
    )
    report["automatic_checks_passed"] = passed
    report["gate"] = (
        "FAILED" if not passed
        else "SIMULATED_ONLY" if provider.metadata.get("simulated", True)
        else "PENDING_SEMANTIC_REVIEW"
    )
    return save()
